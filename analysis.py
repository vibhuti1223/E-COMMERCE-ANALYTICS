# -*- coding: utf-8 -*-
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns

# Load and clean
print("Loading dataset...")
df = pd.read_excel("online_retail_II.xlsx", sheet_name="online_retail_II")
df = df.dropna(subset=["Customer ID", "Description"])
df = df[~df["Invoice"].astype(str).str.startswith("C")]
df = df[df["Quantity"] > 0]
df = df[df["Price"] > 0]
df["TotalRevenue"] = df["Quantity"] * df["Price"]
df["Month"] = df["InvoiceDate"].dt.month
df["Year"] = df["InvoiceDate"].dt.year
print(f"Clean data: {df.shape}")

sns.set_theme(style="whitegrid")

# --- CHART 1: Monthly Revenue Trend ---
monthly_revenue = df.groupby(["Year", "Month"])["TotalRevenue"].sum().reset_index()
monthly_revenue["Period"] = monthly_revenue["Month"].astype(str) + "/" + monthly_revenue["Year"].astype(str)

plt.figure(figsize=(14, 5))
sns.lineplot(data=monthly_revenue, x="Period", y="TotalRevenue", marker="o", color="#1a3c78")
plt.title("Monthly Revenue Trend", fontsize=16, fontweight="bold")
plt.xlabel("Month/Year")
plt.ylabel("Total Revenue (�)")
plt.xticks(rotation=45)
plt.tight_layout()
plt.savefig("chart1_monthly_revenue.png", dpi=150)
plt.close()
print("Chart 1 saved: Monthly Revenue Trend")

# --- CHART 2: Top 10 Countries by Revenue ---
country_revenue = df.groupby("Country")["TotalRevenue"].sum().sort_values(ascending=False).head(10)

plt.figure(figsize=(12, 5))
sns.barplot(x=country_revenue.index, y=country_revenue.values, palette="Blues_d")
plt.title("Top 10 Countries by Revenue", fontsize=16, fontweight="bold")
plt.xlabel("Country")
plt.ylabel("Total Revenue (�)")
plt.xticks(rotation=45)
plt.tight_layout()
plt.savefig("chart2_top_countries.png", dpi=150)
plt.close()
print("Chart 2 saved: Top 10 Countries by Revenue")

# --- CHART 3: Top 10 Best Selling Products ---
top_products = df.groupby("Description")["Quantity"].sum().sort_values(ascending=False).head(10)

plt.figure(figsize=(12, 5))
sns.barplot(x=top_products.values, y=top_products.index, palette="Oranges_r")
plt.title("Top 10 Best Selling Products", fontsize=16, fontweight="bold")
plt.xlabel("Total Quantity Sold")
plt.ylabel("Product")
plt.tight_layout()
plt.savefig("chart3_top_products.png", dpi=150)
plt.close()
print("Chart 3 saved: Top 10 Best Selling Products")

# --- CHART 4: Top 10 Products by Revenue ---
top_revenue_products = df.groupby("Description")["TotalRevenue"].sum().sort_values(ascending=False).head(10)

plt.figure(figsize=(12, 5))
sns.barplot(x=top_revenue_products.values, y=top_revenue_products.index, palette="Greens_r")
plt.title("Top 10 Products by Revenue", fontsize=16, fontweight="bold")
plt.xlabel("Total Revenue (�)")
plt.ylabel("Product")
plt.tight_layout()
plt.savefig("chart4_top_revenue_products.png", dpi=150)
plt.close()
print("Chart 4 saved: Top 10 Products by Revenue")

# --- SUMMARY STATS ---
print("\n===== BUSINESS SUMMARY =====")
print(f"Total Revenue: �{df['TotalRevenue'].sum():,.2f}")
print(f"Total Orders: {df['Invoice'].nunique():,}")
print(f"Total Customers: {df['Customer ID'].nunique():,}")
print(f"Total Products: {df['Description'].nunique():,}")
print(f"Average Order Value: �{df.groupby('Invoice')['TotalRevenue'].sum().mean():,.2f}")
print(f"Top Country: {df.groupby('Country')['TotalRevenue'].sum().idxmax()}")
print(f"Best Month (by revenue): {monthly_revenue.loc[monthly_revenue['TotalRevenue'].idxmax(), 'Period']}")
# --- EXCEL DASHBOARD EXPORT ---
print("\nExporting Excel dashboard...")

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

wb = Workbook()

# --- Sheet 1: Summary ---
ws1 = wb.active
ws1.title = "Summary"

header_fill = PatternFill("solid", fgColor="143C78")
header_font = Font(color="FFFFFF", bold=True, size=12)
title_font = Font(color="143C78", bold=True, size=14)

ws1["A1"] = "E-Commerce Sales Analytics Dashboard"
ws1["A1"].font = Font(color="143C78", bold=True, size=16)
ws1.merge_cells("A1:C1")

ws1["A3"] = "Metric"
ws1["B3"] = "Value"
for cell in [ws1["A3"], ws1["B3"]]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

summary_data = [
    ("Total Revenue (�)", f"�{df['TotalRevenue'].sum():,.2f}"),
    ("Total Orders", f"{df['Invoice'].nunique():,}"),
    ("Total Customers", f"{df['Customer ID'].nunique():,}"),
    ("Total Products", f"{df['Description'].nunique():,}"),
    ("Average Order Value (�)", f"�{df.groupby('Invoice')['TotalRevenue'].sum().mean():,.2f}"),
    ("Top Country", df.groupby("Country")["TotalRevenue"].sum().idxmax()),
    ("Peak Revenue Month", monthly_revenue.loc[monthly_revenue["TotalRevenue"].idxmax(), "Period"]),
]

for i, (metric, value) in enumerate(summary_data, start=4):
    ws1[f"A{i}"] = metric
    ws1[f"B{i}"] = value
    ws1[f"A{i}"].font = Font(bold=True)
    if i % 2 == 0:
        for col in ["A", "B"]:
            ws1[f"{col}{i}"].fill = PatternFill("solid", fgColor="E8EEF7")

ws1.column_dimensions["A"].width = 30
ws1.column_dimensions["B"].width = 25

# --- Sheet 2: Monthly Revenue ---
ws2 = wb.create_sheet("Monthly Revenue")
ws2["A1"] = "Monthly Revenue Trend"
ws2["A1"].font = title_font
ws2.merge_cells("A1:C1")

headers = ["Period", "Year", "Month", "Total Revenue (�)"]
for col, h in enumerate(headers, 1):
    cell = ws2.cell(row=2, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

for i, row in enumerate(monthly_revenue.itertuples(), start=3):
    ws2.cell(row=i, column=1, value=row.Period)
    ws2.cell(row=i, column=2, value=row.Year)
    ws2.cell(row=i, column=3, value=row.Month)
    ws2.cell(row=i, column=4, value=round(row.TotalRevenue, 2))
    if i % 2 == 0:
        for col in range(1, 5):
            ws2.cell(row=i, column=col).fill = PatternFill("solid", fgColor="E8EEF7")

for col in ["A", "B", "C", "D"]:
    ws2.column_dimensions[col].width = 20

# --- Sheet 3: Top Countries ---
ws3 = wb.create_sheet("Top Countries")
ws3["A1"] = "Top 10 Countries by Revenue"
ws3["A1"].font = title_font
ws3.merge_cells("A1:B1")

for col, h in enumerate(["Country", "Total Revenue (�)"], 1):
    cell = ws3.cell(row=2, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

for i, (country, rev) in enumerate(country_revenue.items(), start=3):
    ws3.cell(row=i, column=1, value=country)
    ws3.cell(row=i, column=2, value=round(rev, 2))
    if i % 2 == 0:
        for col in range(1, 3):
            ws3.cell(row=i, column=col).fill = PatternFill("solid", fgColor="E8EEF7")

ws3.column_dimensions["A"].width = 25
ws3.column_dimensions["B"].width = 20

# --- Sheet 4: Top Products ---
ws4 = wb.create_sheet("Top Products")
ws4["A1"] = "Top 10 Products by Revenue"
ws4["A1"].font = title_font
ws4.merge_cells("A1:B1")

for col, h in enumerate(["Product", "Total Revenue (�)"], 1):
    cell = ws4.cell(row=2, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

for i, (product, rev) in enumerate(top_revenue_products.items(), start=3):
    ws4.cell(row=i, column=1, value=product)
    ws4.cell(row=i, column=2, value=round(rev, 2))
    if i % 2 == 0:
        for col in range(1, 3):
            ws4.cell(row=i, column=col).fill = PatternFill("solid", fgColor="E8EEF7")

ws4.column_dimensions["A"].width = 40
ws4.column_dimensions["B"].width = 20

wb.save("ecommerce_dashboard.xlsx")
print("Excel dashboard saved: ecommerce_dashboard.xlsx")